import re
from io import BytesIO
import docx
from fastapi import HTTPException
from PyPDF2 import PdfReader


def extract_pdf_text(file_content: bytes) -> str:
    """Extract text from PDF file content using PyPDF2."""
    try:
        reader = PdfReader(BytesIO(file_content))

        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                raise HTTPException(
                    status_code=400,
                    detail="Password-protected PDF files are not supported."
                )

        text_parts = []
        for page in reader.pages:
            page_text = page.extract_text()
            if page_text and page_text.strip():
                text_parts.append(page_text.strip())

        return "\n".join(text_parts).strip()

    except HTTPException:
        raise
    except Exception:
        raise HTTPException(
            status_code=400,
            detail="Unable to read this PDF. Please upload a valid resume PDF."
        )


def extract_docx_text(file_content: bytes) -> str:
    """Extract text from DOCX file content including paragraphs and tables."""
    try:
        doc = docx.Document(BytesIO(file_content))
        text_parts = []

        for paragraph in doc.paragraphs:
            text = paragraph.text.strip()
            if text:
                text_parts.append(text)

        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    # Deduplicate repeated adjacent cells caused by merged table columns
                    deduped = []
                    for cell_val in row_cells:
                        if not deduped or deduped[-1] != cell_val:
                            deduped.append(cell_val)
                    text_parts.append(" | ".join(deduped))

        return "\n".join(text_parts).strip()

    except HTTPException:
        raise
    except Exception:
        raise HTTPException(
            status_code=400,
            detail="Unable to read this Word document. Please ensure it is a valid DOCX file."
        )


def extract_doc_text(file_content: bytes) -> str:
    """Extract text from legacy .doc files with fallback decoders."""
    # Attempt DOCX parser first (frequently .docx files are saved with .doc extension)
    try:
        docx_result = extract_docx_text(file_content)
        if len(docx_result) >= 50:
            return docx_result
    except Exception:
        pass

    # Fallback: extract printable strings from binary .doc stream
    try:
        decoded = file_content.decode("utf-8", errors="ignore")
        tokens = re.findall(r"[A-Za-z0-9,.:;()/\-\s]{4,}", decoded)
        lines = [token.strip() for token in tokens if len(token.strip()) > 6]
        text = "\n".join(lines).strip()
        if len(text) >= 50:
            return text
    except Exception:
        pass

    raise HTTPException(
        status_code=400,
        detail="Unable to read this legacy .doc file. Please re-save it as .docx or .pdf."
    )


def extract_resume_text(filename: str, file_content: bytes) -> str:
    """Route file to appropriate parser based on extension."""
    if not file_content:
        raise HTTPException(
            status_code=400,
            detail="The uploaded resume file is empty."
        )

    clean_name = (filename or "").lower().strip()

    if clean_name.endswith(".pdf"):
        return extract_pdf_text(file_content)
    elif clean_name.endswith(".docx"):
        return extract_docx_text(file_content)
    elif clean_name.endswith(".doc"):
        return extract_doc_text(file_content)
    else:
        raise HTTPException(
            status_code=400,
            detail="Unsupported file format. Please upload a PDF, DOCX, or DOC resume."
        )
